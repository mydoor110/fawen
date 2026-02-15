"""
电子印章服务
为已审批通过的文档添加电子印章
"""
import os
from datetime import datetime
from typing import Optional
from sqlalchemy.orm import Session
from docx import Document as DocxDocument
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.models.document import Document
from app.models.user import User
from app.models.system import SystemConfig
from app.services.audit_service import log_audit, AuditEvent
from app.utils.logger import get_logger

logger = get_logger("seal_service")


class ElectronicSealService:
    """电子印章服务"""
    
    @staticmethod
    def get_seal_config(db: Session) -> dict:
        """
        获取电子印章配置
        
        Returns:
            配置字典，包含:
            - image_path: 印章图片路径
            - position: 印章位置 ('end'|'custom')
            - width_inches: 印章宽度（英寸）
            - enabled: 是否启用自动盖章
        """
        config = db.query(SystemConfig).filter(
            SystemConfig.config_key == "electronic_seal"
        ).first()
        
        if not config:
            # 返回默认配置
            return {
                "enabled": False,
                "image_path": None,
                "position": "end",
                "width_inches": 1.5
            }
        
        return config.config_value
    
    @staticmethod
    def add_seal_to_document(
        file_path: str, 
        seal_image_path: str,
        width_inches: float = 1.5,
        position: str = "end"
    ) -> bool:
        """
        在Word文档中添加电子印章
        
        Args:
            file_path: Word文档路径
            seal_image_path: 印章图片路径
            width_inches: 印章宽度（英寸）
            position: 印章位置 ('end'表示文档末尾)
            
        Returns:
            是否成功
        """
        try:
            if not os.path.exists(file_path):
                logger.error(f"文档文件不存在: {file_path}")
                return False
            
            if not os.path.exists(seal_image_path):
                logger.error(f"印章图片不存在: {seal_image_path}")
                return False
            
            # 打开Word文档
            doc = DocxDocument(file_path)
            
            if position == "end":
                # 在文档末尾添加印章
                # 添加一个空段落
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT  # 右对齐
                
                # 添加印章图片
                run = paragraph.add_run()
                run.add_picture(seal_image_path, width=Inches(width_inches))
                
                # 添加盖章时间
                timestamp_paragraph = doc.add_paragraph()
                timestamp_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                timestamp_run = timestamp_paragraph.add_run(
                    f"盖章时间: {datetime.now().strftime('%Y年%m月%d日')}"
                )
                timestamp_run.font.size = Pt(10)
            
            # 保存文档
            doc.save(file_path)
            logger.info(f"成功添加电子印章到文档: {file_path}")
            return True
            
        except Exception as e:
            logger.error(f"添加电子印章失败: {str(e)}", exc_info=True)
            return False
    
    @staticmethod
    def seal_document(
        db: Session, 
        document: Document, 
        user: User,
        force: bool = False
    ) -> tuple[bool, str]:
        """
        为文档盖章（主入口函数）
        
        Args:
            db: 数据库会话
            document: 文档对象
            user: 操作用户
            force: 是否强制盖章（忽略已盖章状态）
            
        Returns:
            (是否成功, 消息)
        """
        # 检查文档是否已盖章（可以添加一个字段标记）
        # 这里简化处理，可以通过审计日志判断
        if not force:
            from app.models.system import AuditLog
            existing_seal = db.query(AuditLog).filter(
                AuditLog.resource_type == "Document",
                AuditLog.resource_id == str(document.id),
                AuditLog.action == AuditEvent.DOCUMENT_SEAL
            ).first()
            
            if existing_seal:
                return False, "文档已盖章，如需重新盖章请使用强制盖章功能"
        
        # 检查文档文件是否存在
        if not document.file_path or not os.path.exists(document.file_path):
            return False, "文档文件不存在或路径无效"
        
        # 获取印章配置
        seal_config = ElectronicSealService.get_seal_config(db)
        
        if not seal_config.get("enabled"):
            logger.warning("电子印章功能未启用")
            return False, "电子印章功能未启用，请联系系统管理员配置"
        
        seal_image_path = seal_config.get("image_path")
        if not seal_image_path:
            return False, "未配置印章图片路径"
        
        # 执行盖章
        success = ElectronicSealService.add_seal_to_document(
            document.file_path,
            seal_image_path,
            width_inches=seal_config.get("width_inches", 1.5),
            position=seal_config.get("position", "end")
        )
        
        if not success:
            return False, "盖章失败，请检查文档和印章图片"
        
        # 记录审计日志
        log_audit(db, user, AuditEvent.DOCUMENT_SEAL,
                  "Document", str(document.id),
                  {
                      "seal_time": datetime.utcnow().isoformat(),
                      "seal_image": seal_image_path,
                      "force": force
                  })
        
        db.commit()
        
        logger.info(f"文档 {document.id} 盖章成功，操作人: {user.username}")
        return True, "盖章成功"


# 便捷函数
def seal_document(db: Session, document: Document, user: User, force: bool = False) -> tuple[bool, str]:
    """
    便捷函数：为文档盖章
    
    Args:
        db: 数据库会话
        document: 文档对象
        user: 操作用户
        force: 是否强制盖章
        
    Returns:
        (是否成功, 消息)
        
    Examples:
        >>> success, msg = seal_document(db, document, current_user)
        >>> if success:
        ...     print("盖章成功")
        ... else:
        ...     print(f"盖章失败: {msg}")
    """
    return ElectronicSealService.seal_document(db, document, user, force)
